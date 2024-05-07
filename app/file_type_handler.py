import csv
import io
import os
from fastapi import HTTPException
import openpyxl
import pinecone
from dotenv import load_dotenv
from langchain_openai import OpenAIEmbeddings
from langchain_community.vectorstores import Pinecone
from langchain_community.document_loaders import PyPDFLoader
from langchain.document_loaders.csv_loader import CSVLoader
from langchain_text_splitters import CharacterTextSplitter
from langchain.schema import Document

load_dotenv()
api_key = os.environ.get('OPENAI_API_KEY')
index_name = os.environ.get('PINECONE_API_INDEX')
if not api_key:
    print('OpenAI API key not found in environment variables.')
    exit()
# initialize pinecone
pinecone.init(
    api_key=os.environ.get("PINECONE_API_KEY"),
    environment = os.environ.get("PINECONE_API_ENV")
)
# initialize embedding function
embedding_function = OpenAIEmbeddings()

def find_delimiter(first_line):
    delimiters = [',', ';', '\t', '|']
    for delimiter in delimiters:
        if delimiter in first_line:
            return delimiter
    return ","  # Default to comma if no delimiter found

def upload_to_namespace(namespace, data, file_path):
    index = pinecone.Index(index_name)
    vectorstore = Pinecone(index, embedding=embedding_function, text_key = 'csv_text', namespace = namespace)
    vectorstore.add_documents(documents = data)
    os.remove(file_path)

def add_pdf_document(namespace, file_path):
    loader = PyPDFLoader(file_path)
    pages = loader.load_and_split()
    upload_to_namespace(namespace, pages, file_path)
    
def add_csv_documents(namespace, file_path, delimiter):
    loader = CSVLoader(file_path=file_path, encoding='utf-8', csv_args={'delimiter': delimiter})
    data = loader.load()
    upload_to_namespace(namespace, data, file_path)
    
def add_txt_documents(namespace, file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        txt_data = f.read()
    separator = '\n'
    text_splitter = CharacterTextSplitter(
        separator = separator,
        chunk_size = 2000,
        chunk_overlap = 0,
        length_function = len,
        is_separator_regex = False
    )
    texts = text_splitter.create_documents([txt_data])
    c=0
    new_docs = []
    for text in texts:
        page_content = text.page_content
        metadata = {'document': 'abc'}
        new_doc = Document(page_content, metadata=metadata)
        new_docs.append(new_doc)
        c+=1
    upload_to_namespace(namespace, new_docs, file_path)
    
def read_the_structured_file(filepath):
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"The file {filepath} does not exist.")
    filename = os.path.basename(filepath)
    if filename.endswith('.csv'):
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                contents_str = f.read()
        except UnicodeDecodeError as e:
            raise HTTPException(status_code=400, detail="Unable to decode contents: " + str(e))
        first_line = contents_str.split('\n')[0]
        try:
            delimiter = find_delimiter(first_line)
        except Exception as e:
            raise HTTPException(status_code=400, detail="Unable to find delimiter: " + str(e))
        try:
            reader = csv.DictReader(io.StringIO(contents_str), delimiter=delimiter)
        except csv.Error as e:
            raise HTTPException(status_code=400, detail="Unable to read CSV: " + str(e))
        return reader, delimiter
    elif filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filename=filepath)
        sheet = wb.active
        output = io.StringIO()
        csv_file = csv.writer(output)
        for row in sheet.iter_rows(values_only=True):
            csv_file.writerow(row)
        csv_data = output.getvalue()
        csv_file = io.StringIO(csv_data)
        first_line = csv_data.split('\n')[0]
        delimiter = find_delimiter(first_line)
        reader = csv.DictReader(csv_file)
        return reader, delimiter
    

    